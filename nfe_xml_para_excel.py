import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NS = {"nfe": "http://www.portalfiscal.inf.br/nfe"}


def txt(no, caminho, padrao=""):
    """Retorna o texto de um sub-elemento (ou padrão se não existir)."""
    if no is None:
        return padrao
    el = no.find(caminho, NS)
    return el.text.strip() if el is not None and el.text else padrao


def num(valor):
    """Converte string numérica do XML (ponto decimal) em float."""
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def ler_nfe(caminho_xml):
    """Lê um XML de NF-e e devolve (dados_da_nota, lista_de_itens)."""
    raiz = ET.parse(caminho_xml).getroot()

    # Funciona tanto com <nfeProc> (nota autorizada) quanto com <NFe> solta
    inf = raiz.find(".//nfe:infNFe", NS)
    if inf is None:
        raise ValueError("Não é um XML de NF-e válido (infNFe não encontrado)")

    chave = inf.get("Id", "").replace("NFe", "")
    ide = inf.find("nfe:ide", NS)
    emit = inf.find("nfe:emit", NS)
    dest = inf.find("nfe:dest", NS)
    tot = inf.find("nfe:total/nfe:ICMSTot", NS)

    nota = {
        "Chave": chave,
        "Número": txt(ide, "nfe:nNF"),
        "Série": txt(ide, "nfe:serie"),
        "Data Emissão": txt(ide, "nfe:dhEmi")[:10] or txt(ide, "nfe:dEmi"),
        "Natureza Operação": txt(ide, "nfe:natOp"),
        "Emitente CNPJ": txt(emit, "nfe:CNPJ"),
        "Emitente Nome": txt(emit, "nfe:xNome"),
        "Emitente UF": txt(emit, "nfe:enderEmit/nfe:UF"),
        "Destinatário CNPJ/CPF": txt(dest, "nfe:CNPJ") or txt(dest, "nfe:CPF"),
        "Destinatário Nome": txt(dest, "nfe:xNome"),
        "Destinatário UF": txt(dest, "nfe:enderDest/nfe:UF"),
        "Valor Produtos": num(txt(tot, "nfe:vProd")),
        "Frete": num(txt(tot, "nfe:vFrete")),
        "Desconto": num(txt(tot, "nfe:vDesc")),
        "ICMS": num(txt(tot, "nfe:vICMS")),
        "IPI": num(txt(tot, "nfe:vIPI")),
        "Valor Total NF": num(txt(tot, "nfe:vNF")),
        "Arquivo": Path(caminho_xml).name,
    }

    itens = []
    for det in inf.findall("nfe:det", NS):
        prod = det.find("nfe:prod", NS)
        itens.append({
            "Chave": chave,
            "Número NF": nota["Número"],
            "Item": det.get("nItem", ""),
            "Código": txt(prod, "nfe:cProd"),
            "Descrição": txt(prod, "nfe:xProd"),
            "NCM": txt(prod, "nfe:NCM"),
            "CFOP": txt(prod, "nfe:CFOP"),
            "Unidade": txt(prod, "nfe:uCom"),
            "Quantidade": num(txt(prod, "nfe:qCom")),
            "Valor Unitário": num(txt(prod, "nfe:vUnCom")),
            "Valor Total": num(txt(prod, "nfe:vProd")),
        })

    return nota, itens


def escrever_aba(ws, linhas):
    """Escreve uma lista de dicts numa aba, com formatação básica."""
    if not linhas:
        return
    cabecalho = list(linhas[0].keys())
    ws.append(cabecalho)
    for linha in linhas:
        ws.append([linha[c] for c in cabecalho])

    # Cabeçalho
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="305496")
        cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Largura das colunas + formato numérico
    for i, nome in enumerate(cabecalho, start=1):
        maior = max(len(str(c.value)) if c.value is not None else 0
                    for c in ws[get_column_letter(i)])
        ws.column_dimensions[get_column_letter(i)].width = min(maior + 2, 50)
        if any(k in nome for k in ("Valor", "Frete", "Desconto", "ICMS", "IPI")):
            for c in ws[get_column_letter(i)][1:]:
                c.number_format = "#,##0.00"
